"""Append AI-arbitration failures to 中国执照库/待人工处理.xlsx."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

HEADERS = ("时间", "文件夹", "文件", "原因", "名称", "信用代码", "缺失字段")
SHEET_NAME = "待人工处理"
FILE_NAME = "待人工处理.xlsx"


def review_sheet_path(source_dir: str | Path) -> Path:
    return Path(source_dir) / FILE_NAME


def append_review_row(
    source_dir: str | Path,
    *,
    folder: str = "",
    file: str = "",
    reason: str = "",
    name: str = "",
    credit_code: str = "",
    missing: str = "",
) -> Path:
    """Create the workbook if missing, then append one failure row."""
    import openpyxl
    from openpyxl.styles import Font

    path = review_sheet_path(source_dir)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = openpyxl.load_workbook(path)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(list(HEADERS))
        for cell in ws[1]:
            cell.font = Font(bold=True)
    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        folder,
        file,
        reason,
        name,
        credit_code,
        missing,
    ])
    wb.save(path)
    wb.close()
    return path
